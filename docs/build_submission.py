from pathlib import Path
from datetime import date
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak,
    Table, TableStyle, KeepTogether, ListFlowable, ListItem
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "Tech_Challenge-Fase1-Grupo76-completo.docx"
OUT_PDF = ROOT / "docs" / "Tech_Challenge-Fase1-Grupo76-entrega.pdf"

BLUE = "19324D"
MID_BLUE = "2E5D7B"
LIGHT_BLUE = "EAF1F5"
PALE_BLUE = "F4F7F9"
GOLD = "B8831B"
PALE_GOLD = "FFF4CC"
GREEN = "217346"
PALE_GREEN = "EAF5EE"
RED = "9B1C1C"
PALE_RED = "FBEAEA"
GRAY = "5E6872"
LIGHT_GRAY = "E5E9ED"
WHITE = "FFFFFF"

REPO = "https://github.com/kayleighszs/oficinatrack"
BRANCH = "master"
COMMIT = "02ce391039d5e3616c669e894c05a0fb864416da"
COMMIT_URL = f"{REPO}/tree/{COMMIT}"

MIRO = {
    "Event Storming": "https://miro.com/welcomeonboard/cEpNN1pQTGxqSzJieDEwLzN3ZjRmNEZWM2xiOW00S0tZOVBJaUtwWVozZFZIbHQwY3ZhWGNGZE1sU0tVakJYTzV4U2w4L0F6T0x0bFdkdnFhVXMvUU9ac3Iwc1RHVEJITUtyOTRFWUp1bG9xT0RZeUs3bUxERS9jWHlzdzBCTWtNakdSWkpBejJWRjJhRnhhb1UwcS9BPT0hdjE=?share_link_id=89439599795",
    "Bounded Contexts": "https://miro.com/welcomeonboard/Y0pVSFJzV2NaUHhGSjExaks2S2FEelptVHFBQ0lXMFQrbmVBSGF2SWVuWkVXditBV2o4U0F3dHVyT0hndVlRWXdHRHVHMFdMOG12cXpoZEZDeHVmMDZmblptTUx0VHBtRml6bGd6bjNnbE8xS2EyaWFaY3UrZnhUNzJFRjk4M0JhWWluRVAxeXRuUUgwWDl3Mk1qRGVRPT0hdjE=?share_link_id=835649372749",
    "Domain Storytelling": "https://miro.com/welcomeonboard/dU1scWJkWVFXUitoU1lWVlNua3B5bk9qWldGOGJyelg4ZUdYMWNyR2lLb1NXTGEybHlYTmpFZVhkaFdxcWFtcGN4cmN2VjFLNVFoenZQblJGVEQyTzZmblptTUx0VHBtRml6bGd6bjNnbE1QT1QyOUQ1V05ISE5LTWJoaThjTWZ0R2lncW1vRmFBVnlLcVJzTmdFdlNRPT0hdjE=?share_link_id=454382467708",
    "Tactical Design": "https://miro.com/welcomeonboard/SWpmRTdqWk1takdKaFNEaFB6dEZONTM5VHMrdUVwMSs3cGJpcHNYaFVHVitURW5rbmRuWWxueTR3UzJPaTYxRUl4ZVlKWDdEVSs3dFB6WS92N3FkSGFmblptTUx0VHBtRml6bGd6bjNnbE96WkQ0Mk1maDJTNURNSmFPd2svTkRBd044SHFHaVlWYWk0d3NxeHNmeG9BPT0hdjE=?share_link_id=267198620746",
}

MEMBERS = [
    ("Allan de Lima Carvalho Silva", "RM 378457", "zonahk"),
    ("Guilherme Azevedo da Trindade", "RM 378112", "guilhermetrindade97_75307"),
    ("Laura Henrique Alves", "RM 377327", "maflexninja"),
    ("Lucas de Souza Cancela", "RM 377972", "lucancela"),
    ("PENDENTE — nome completo de Kayleigh", "PENDENTE — RM", "kayleighsz_"),
]

TECH = (
    "Java 21; Spring Boot 4.1.1; Spring Web MVC; Spring Security e JWT; "
    "Spring Data JPA/Hibernate; PostgreSQL; Flyway; OpenAPI/Swagger; Docker e "
    "Docker Compose; JUnit 5; Mockito; Testcontainers; JaCoCo."
)

COMPLIANCE = [
    ("Identificação da entrega", "Curso, fase, grupo, integrantes, RMs, Discord e stack na capa.", "ATENÇÃO", "Nome completo e RM de Kayleigh pendentes."),
    ("Repositório público e versionado", f"{REPO} — branch {BRANCH}; commit {COMMIT}.", "VALIDAR", "Confirmar visibilidade pública e proteção da branch principal durante correção."),
    ("README e execução local", "README, Dockerfile, docker-compose.yml e .env.example.", "PARCIAL", "README funciona, mas deve receber links do DDD/vídeo e seção arquitetural mais completa."),
    ("Vídeo público no YouTube", "Demonstração de até 15 minutos vinculada ao commit entregue.", "PENDENTE", "Inserir URL pública antes de exportar versão final."),
    ("Documentação DDD", "Quatro quadros Miro e três diagramas Excalidraw no diretório docs/.", "VALIDAR", "Confirmar acesso anônimo aos quadros; ideal exportar também para PDF/PNG em docs/ddd/."),
    ("Funcionalidades obrigatórias", "CRUDs, OS, orçamento, consulta pública, estoque, relatório e histórico presentes.", "ATENDIDO", "Evidência mapeada a pacotes e endpoints neste documento."),
    ("JWT e validações", "Rotas administrativas autenticadas; roles; CPF/CNPJ e placa validados.", "ATENDIDO", "JWT HS256, BCrypt, sessão stateless e token público por OS com hash SHA-256."),
    ("Testes e cobertura", "./mvnw verify: 179 testes; 88,41% de cobertura de linhas no escopo JaCoCo.", "ATENDIDO", "Execução validada em 20/10/2026 com PostgreSQL via Testcontainers."),
    ("Análise de vulnerabilidades", "Nenhum relatório versionado foi localizado.", "PENDENTE", "Executar SCA, revisar achados e incluir relatório/resultado no repositório."),
    ("PDF único", "Este arquivo reúne identificação, links, decisões e evidências.", "ATENDIDO", "Substituir pendências antes do envio oficial."),
]

CONTEXTS = [
    ("Ordens de Serviço", "Core Domain", "OrdemServico", "Orcamento, itens e histórico", "Coordena diagnóstico, orçamento, decisão, execução e entrega."),
    ("Clientes", "Apoio", "Cliente", "TipoPessoa; CPF/CNPJ validado", "Mantém responsável pelo atendimento e estado ativo."),
    ("Veículos", "Apoio", "Veiculo", "Placa", "Vincula veículo ao cliente e normaliza placa."),
    ("Catálogo de Serviços", "Apoio", "Servico", "valor e tempo estimado", "Fonte dos serviços e preços adicionados à OS."),
    ("Estoque", "Apoio", "Peca", "MovimentacaoEstoque", "Controla saldo, mínimo e trilha de entradas/saídas."),
    ("Acesso Público", "Apoio", "AcessoPublicoOrdemServico", "token opaco e validade", "Consulta e decisão do cliente sem expor JWT administrativo."),
    ("Identidade e Acesso", "Genérico", "Usuario", "Role e JWT", "Autentica equipe e aplica autorização por perfil."),
    ("Relatórios", "Apoio", "Consulta", "TempoMedioExecucaoResponse", "Calcula média entre EM_EXECUCAO e FINALIZADA."),
]

DDD_ELEMENTS = [
    ("Agregado OS", "OrdemServico (raiz); Orcamento; ItemServicoOrdem; ItemPecaOrdem; HistoricoStatusOrdemServico", "Mantém transições, versões de orçamento, composição e histórico consistentes."),
    ("Agregado Estoque", "Peca (raiz); MovimentacaoEstoque", "Impede saldo negativo e registra origem manual ou automática."),
    ("Entidades", "Cliente, Veiculo, Servico, Usuario, AcessoPublicoOrdemServico", "Possuem identidade e ciclo de vida persistido."),
    ("Objeto de valor", "Placa", "Normaliza formato antigo/Mercosul e rejeita placa inválida."),
    ("Conceitos enumerados", "StatusOrdemServico, StatusOrcamento, TipoPessoa, Role, TipoMovimentacao, Origem", "Vocabulário fechado usado por regras e persistência."),
    ("Repositórios", "Interfaces Spring Data por módulo", "Abstraem persistência; consultas com lock protegem decisão/estoque concorrente."),
    ("Serviços de aplicação", "*Service em cada contexto", "Orquestram repositórios, transações e chamadas aos agregados."),
]

GLOSSARY = [
    ("Ordem de Serviço (OS)", "Atendimento de um veículo, do recebimento à entrega."),
    ("Diagnóstico", "Etapa em que serviços e peças necessários são identificados."),
    ("Orçamento", "Versão calculada com subtotais de serviços e peças, submetida ao cliente."),
    ("Acesso público", "Credencial opaca, temporária e vinculada a uma única OS."),
    ("Baixa automática", "Saída de estoque feita na mesma transação da aprovação do orçamento."),
    ("Histórico de status", "Registro imutável das transições e respectivos instantes."),
    ("Peça/insumo", "Material consumível controlado por saldo e estoque mínimo."),
    ("Tempo médio de execução", "Média entre entrada em EM_EXECUCAO e entrada em FINALIZADA."),
]

EVENT_FLOWS = [
    ("Criação e orçamento da OS", "Atendente / mecânico", "Cadastrar cliente e veículo; abrir OS; iniciar diagnóstico; incluir serviço/peça; calcular e enviar orçamento", "ClienteIdentificado; VeiculoCadastrado; OrdemServicoCriada; DiagnosticoIniciado; OrcamentoCalculado; OrcamentoEnviado", "Veículo deve pertencer ao cliente; itens só entram em diagnóstico; orçamento exige ao menos um item; somente versão atual pode ser enviada."),
    ("Acompanhamento e decisão", "Cliente / atendente / mecânico", "Emitir acesso; consultar status/orçamento; aprovar ou recusar; finalizar; entregar", "AcessoPublicoEmitido; OrcamentoAprovado ou OrcamentoRecusado; StatusAlterado; OrdemFinalizada; VeiculoEntregue", "Token deve estar válido e vinculado à OS; decisão é idempotente; transições seguem máquina de estados."),
    ("Gestão de peças e insumos", "Gestor / sistema", "Cadastrar peça; registrar entrada/saída; aprovar orçamento", "EstoqueMovimentado; EstoqueBaixado", "Quantidade positiva; saldo não pode ficar negativo; aprovação bloqueia linhas para atualização e registra OS/orçamento na movimentação."),
]

ENDPOINTS = [
    ("Autenticação", "POST /auth/login", "Público", "Valida credenciais e emite JWT."),
    ("Clientes", "POST /clientes; GET /clientes; GET /clientes/{id}; PUT /clientes/{id}; PATCH /clientes/{id}/inativar", "JWT", "Cadastro, busca, atualização e inativação."),
    ("Veículos", "POST /veiculos; GET /veiculos; GET /veiculos/{id}; PATCH /veiculos/{id}", "JWT", "Cadastro vinculado ao cliente, consulta e atualização/inativação."),
    ("Serviços", "POST /servicos; GET /servicos; GET /servicos/{id}; PUT /servicos/{id}; PATCH /servicos/{id}/inativar", "JWT", "Catálogo de serviços."),
    ("Peças", "POST /pecas; GET /pecas; GET /pecas/{id}; PUT /pecas/{id}; PATCH /pecas/{id}/estoque", "JWT", "Cadastro, filtros e movimentação manual."),
    ("OS — fluxo", "POST /ordens-servico; POST /{id}/diagnosticos; POST /{id}/itens-servico; POST /{id}/itens-peca; POST /{id}/orcamentos; POST /{id}/orcamentos/{orcamentoId}/envios; POST /{id}/finalizacoes; POST /{id}/entregas", "JWT", "Fluxo administrativo e transições."),
    ("OS — consultas", "GET /ordens-servico; GET /ordens-servico/{id}; GET /ordens-servico/{id}/historico-status", "JWT", "Lista paginada com filtros, detalhe e histórico."),
    ("Acesso da OS", "POST /ordens-servico/{ordemId}/acessos-publicos; DELETE /ordens-servico/{ordemId}/acessos-publicos/{acessoId}", "JWT", "Emite token uma vez e permite revogação."),
    ("Cliente", "GET /publico/ordens-servico/{numeroOs}/status; GET /publico/ordens-servico/{numeroOs}/orcamentos/{orcamentoId}; POST /publico/ordens-servico/{numeroOs}/orcamentos/{orcamentoId}/decisoes", "Token da OS", "Consulta e decisão do orçamento."),
    ("Relatórios", "GET /relatorios/tempo-medio-execucao", "JWT/GESTOR", "Média do tempo de execução."),
]

REFERENCES = [
    ("Enunciado oficial da fase", "docs/Fase 1 - Tech Challenge.pdf"),
    ("Domain-Driven Design", "https://martinfowler.com/bliki/DomainDrivenDesign.html"),
    ("Bounded Context", "https://martinfowler.com/bliki/BoundedContext.html"),
    ("Spring Security — JWT", "https://docs.spring.io/spring-security/reference/servlet/oauth2/resource-server/jwt.html"),
    ("PostgreSQL — Transactions", "https://www.postgresql.org/docs/current/tutorial-transactions.html"),
    ("OWASP Dependency-Check", "https://owasp.org/projects/dependency-check"),
]

MAIN_STEPS = [
    "Atendente autentica e recebe JWT.",
    "Cliente e veículo são cadastrados; veículo deve pertencer ao cliente.",
    "OS é aberta em RECEBIDA; acesso público pode ser emitido e entregue ao cliente.",
    "Mecânico inicia diagnóstico; OS passa a EM_DIAGNOSTICO.",
    "Serviços e peças são incluídos com quantidade positiva e snapshot de preço.",
    "Sistema calcula versão do orçamento e soma subtotais com escala monetária de duas casas.",
    "Atendente envia apenas o orçamento atual; OS passa a AGUARDANDO_APROVACAO.",
    "Cliente consulta e aprova/recusa usando token opaco vinculado à OS.",
    "Na aprovação, peças são bloqueadas e debitadas; movimentações são registradas; OS passa a EM_EXECUCAO.",
    "Mecânico finaliza; atendente confirma entrega; histórico registra todas as mudanças.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def hyperlink(paragraph, text, url, color=MID_BLUE):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string("1E2933")
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after, color in [
        ("Title", 25, 0, 10, BLUE), ("Subtitle", 13, 0, 8, GRAY),
        ("Heading 1", 16, 16, 8, BLUE), ("Heading 2", 13, 12, 6, MID_BLUE),
        ("Heading 3", 11.5, 8, 4, BLUE),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = name != "Subtitle"
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def doc_para(doc, text="", style=None, bold_prefix=None, align=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    if align is not None: p.alignment = align
    return p


def doc_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65); p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3); p.add_run(item)


def doc_table(doc, headers, rows, widths=None, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h; set_cell_shading(hdr[i], BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if status_col == i:
                fill = PALE_GREEN if value == "ATENDIDO" else PALE_RED if value == "PENDENTE" else PALE_GOLD
                set_cell_shading(cells[i], fill)
    for ri, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell); set_cell_border(cell)
            if widths: cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = "Calibri"; r.font.size = Pt(8.5 if len(headers) >= 4 else 9)
                    if ri == 0: r.font.bold = True; r.font.color.rgb = RGBColor.from_string(WHITE)
                    elif status_col == i: r.font.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def doc_callout(doc, title, text, kind="info"):
    fill, accent = (PALE_RED, RED) if kind == "danger" else (PALE_GOLD, GOLD) if kind == "warn" else (LIGHT_BLUE, MID_BLUE)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    cell = t.cell(0, 0); set_cell_shading(cell, fill); set_cell_margins(cell, 140, 180, 140, 180); set_cell_border(cell, accent, "10")
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph(text); p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def doc_link_line(doc, label, url):
    p = doc.add_paragraph(); p.add_run(label + ": ").bold = True; hyperlink(p, url, url); return p


def build_docx():
    doc = Document(); style_doc(doc)
    core = doc.core_properties
    core.title = "Tech Challenge — Fase 1 — OficinaTrack — Grupo 76"
    core.subject = "Documentação de entrega do MVP de backend"
    core.author = "Grupo 76"
    core.keywords = "FIAP, SOAT, DDD, Java, Spring Boot, OficinaTrack"

    # Cover
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(26)
    r = p.add_run("FIAP"); r.bold = True; r.font.size = Pt(34); r.font.color.rgb = RGBColor.from_string(RED)
    doc_para(doc, "SOAT — Software Architecture", "Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc_para(doc, "TECH CHALLENGE — FASE 1", "Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc_para(doc, "OFICINATRACK", align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].font.size = Pt(22); p.runs[0].bold = True; p.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    doc_para(doc, "Sistema Integrado de Atendimento e Execução de Serviços", "Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc_para(doc, "GRUPO 76", align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold = True; p.runs[0].font.size = Pt(14)
    doc_table(doc, ["Integrante", "RM", "Discord"], MEMBERS, [8.4, 3.0, 5.0])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Linguagem e tecnologias\n").bold = True; p.add_run(TECH)
    doc_para(doc, "São Paulo — 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # Control and index
    doc_para(doc, "Controle da entrega", "Heading 1")
    doc_table(doc, ["Campo", "Valor"], [
        ("Curso / fase / grupo", "SOAT — Software Architecture / Fase 1 / Grupo 76"),
        ("Repositório", REPO), ("Branch de entrega", BRANCH), ("Commit de referência", COMMIT),
        ("Versão do documento", "1.0 — 20/10/2026"),
        ("Vídeo", "PENDENTE — inserir URL pública do YouTube"),
    ], [4.2, 12.2])
    doc_link_line(doc, "Repositório completo", REPO)
    doc_link_line(doc, "Versão congelada pelo commit", COMMIT_URL)
    doc_callout(doc, "Bloqueios antes da entrega", "Preencher nome completo e RM de Kayleigh; inserir link público do vídeo; confirmar Miro sem login; anexar resultado da análise de vulnerabilidades; confirmar branch principal protegida e versão final via PR.", "danger")
    doc_para(doc, "Sumário", "Heading 1")
    doc_bullets(doc, [
        "1. Visão executiva", "2. Matriz de aderência ao enunciado", "3. Arquitetura e decisões técnicas",
        "4. Domain-Driven Design", "5. Fluxos e regras de negócio", "6. APIs e segurança",
        "7. Persistência e transações", "8. Qualidade e evidências", "9. Execução e demonstração",
        "10. Links DDD, roteiro do vídeo e checklist final", "Referências",
    ])
    doc.add_page_break()

    doc_para(doc, "1. Visão executiva", "Heading 1")
    doc_para(doc, "O OficinaTrack é um MVP de back-end monolítico para substituir anotações e planilhas no atendimento de uma oficina mecânica. A solução concentra cadastros, ordens de serviço, diagnóstico, orçamento, autorização do cliente, consumo de peças, histórico e indicadores em uma API REST documentada.")
    doc_para(doc, "Problema e resultado esperado", "Heading 2")
    doc_bullets(doc, [
        "Reduzir erros de priorização e dar rastreabilidade ao status da OS.",
        "Evitar perda de histórico de clientes, veículos, orçamentos e serviços.",
        "Controlar peças e insumos com movimentações auditáveis e baixa atômica.",
        "Permitir aprovação do orçamento pelo cliente sem acesso administrativo.",
        "Proteger operações internas com JWT e perfis de acesso.",
    ])
    doc_para(doc, "Escopo entregue", "Heading 2")
    doc_table(doc, ["Dimensão", "Evidência"], [
        ("Aplicação", "88 arquivos Java organizados em nove módulos de negócio/técnicos; 37 endpoints REST."),
        ("Banco", "PostgreSQL com quatro migrations Flyway e constraints de integridade."),
        ("Qualidade", "179 testes em 19 classes; build Maven validado; cobertura de linhas de 88,41%."),
        ("Operação", "Dockerfile, Docker Compose, variáveis de ambiente e Swagger UI."),
        ("DDD", "Diagramas de domínio, AS IS/TO BE e épicos; quatro quadros Miro; mapeamento código-domínio neste PDF."),
    ], [3.2, 13.2])

    doc_para(doc, "2. Matriz de aderência ao enunciado", "Heading 1")
    doc_para(doc, "Status baseado no código e nos artefatos locais em 20/10/2026. Itens pendentes não são apresentados como concluídos.")
    doc_table(doc, ["Requisito", "Evidência", "Status", "Ação/observação"], COMPLIANCE, [3.4, 6.0, 2.2, 4.8], 2)

    doc_para(doc, "3. Arquitetura e decisões técnicas", "Heading 1")
    doc_para(doc, "3.1 Estilo arquitetural", "Heading 2")
    doc_para(doc, "Monólito modular em camadas. Cada contexto reúne API, aplicação, domínio e infraestrutura. Controllers tratam HTTP; serviços de aplicação demarcam casos de uso e transações; entidades concentram invariantes; repositórios Spring Data isolam persistência. O shared contém validações, tratamento de erros e configuração OpenAPI.")
    doc_callout(doc, "Fluxo de uma requisição", "Cliente HTTP → Controller/DTO → Service transacional → Entidade/agregado → Repository → PostgreSQL. Respostas e erros retornam contratos próprios, sem expor entidades JPA diretamente.")
    doc_para(doc, "3.2 Módulos", "Heading 2")
    doc_table(doc, ["Pacote", "Responsabilidade", "Camadas principais"], [
        ("auth", "Login, JWT, usuários e roles", "api / application / domain / infra"),
        ("cliente", "Cadastro e validação de documento", "api / application / domain / infra / validation"),
        ("veiculo", "Veículo, vínculo com cliente e placa", "api / application / domain / infra"),
        ("catalogoservicos", "Serviços, preço e tempo estimado", "api / application / domain / infra"),
        ("estoque", "Peças e movimentações", "api / application / domain / infrastructure"),
        ("ordemservico", "Core Domain: OS, itens, orçamento e histórico", "api / application / domain / infra"),
        ("publico", "Credencial por OS, consulta e decisão", "api / application / domain / infra"),
        ("relatorios", "Indicador de tempo médio", "api / application"),
        ("shared", "Erros, configuração e validações comuns", "config / exception / validation"),
    ], [3.3, 7.0, 6.1])
    doc_para(doc, "3.3 Decisão de banco de dados", "Heading 2")
    doc_para(doc, "PostgreSQL foi escolhido por oferecer transações ACID, constraints, locks de linha e consultas relacionais adequadas ao fluxo. OS, itens, versões de orçamento, histórico e movimentações possuem relações fortes; consistência é mais importante que escalabilidade independente no MVP. Flyway mantém evolução reprodutível do schema. A aprovação carrega a OS e bloqueia as peças em ordem determinística, reduzindo risco de corrida e deadlock.")
    doc_para(doc, "3.4 Tecnologias e justificativas", "Heading 2")
    doc_table(doc, ["Tecnologia", "Papel"], [
        ("Java 21 / Spring Boot 4.1.1", "Base da API e injeção de dependências."),
        ("Spring Web MVC", "Controllers REST, validação de entrada e serialização JSON."),
        ("Spring Data JPA / Hibernate", "Mapeamento relacional, repositories e locking."),
        ("Spring Security / JJWT", "Autenticação stateless, BCrypt e JWT HS256."),
        ("PostgreSQL / Flyway", "Persistência transacional e migrations versionadas."),
        ("springdoc OpenAPI", "Swagger UI e contrato navegável."),
        ("Docker / Compose", "Ambiente reprodutível para API e banco."),
        ("JUnit / Mockito / Testcontainers / JaCoCo", "Testes isolados, integração real e cobertura mínima."),
    ], [5.0, 11.4])

    doc_para(doc, "4. Domain-Driven Design", "Heading 1")
    doc_para(doc, "4.1 Estratégia e Bounded Contexts", "Heading 2")
    doc_para(doc, "O domínio central é a evolução da Ordem de Serviço. Cadastros, estoque, identidade e relatórios dão suporte. No monólito, limites são módulos de código e linguagem, não serviços distribuídos.")
    doc_table(doc, ["Contexto", "Tipo", "Raiz/centro", "Elementos", "Responsabilidade"], CONTEXTS, [3.0, 2.0, 3.0, 3.4, 5.0])
    doc_para(doc, "Relações do Context Map", "Heading 3")
    doc_bullets(doc, [
        "Clientes → Veículos: cliente é fornecedor da identidade usada no vínculo do veículo.",
        "Clientes e Veículos → OS: a aplicação valida existência, atividade e pertencimento antes de abrir atendimento.",
        "Catálogo e Estoque → OS: nomes e preços são copiados para itens, preservando o orçamento histórico.",
        "OS → Estoque: aprovação aciona baixa síncrona e transacional; estoque impede saldo negativo.",
        "OS → Acesso Público: token limita consulta e decisão a uma OS específica.",
        "Identidade e Acesso é transversal às rotas administrativas; Relatórios lê histórico produzido pela OS.",
    ])
    doc_para(doc, "4.2 Modelo tático e relação com o código", "Heading 2")
    doc_table(doc, ["Categoria", "Implementação", "Regra principal"], DDD_ELEMENTS, [3.0, 7.2, 6.2])
    doc_callout(doc, "Nota de fidelidade", "Os eventos abaixo descrevem fatos de negócio usados no Event Storming. Na implementação atual, a maioria aparece como mudança de estado/histórico dentro da transação, não como classes publicadas em barramento. A baixa de estoque também é síncrona no serviço de consulta pública.", "warn")
    doc_para(doc, "4.3 Linguagem Ubíqua", "Heading 2")
    doc_table(doc, ["Termo", "Definição adotada"], GLOSSARY, [5.0, 11.4])
    doc_para(doc, "4.4 Event Storming consolidado", "Heading 2")
    doc_table(doc, ["Fluxo", "Atores", "Comandos", "Eventos/fatos", "Políticas"], EVENT_FLOWS, [2.8, 2.5, 4.3, 4.3, 4.5])

    doc_para(doc, "5. Fluxos e regras de negócio", "Heading 1")
    doc_para(doc, "5.1 Fluxo principal da Ordem de Serviço", "Heading 2")
    for i, s in enumerate(MAIN_STEPS, 1): doc_para(doc, f"{i}. {s}")
    doc_para(doc, "5.2 Máquina de estados", "Heading 2")
    doc_table(doc, ["Origem", "Comando", "Destino", "Proteção"], [
        ("—", "Abrir OS", "RECEBIDA", "Cliente/veículo válidos"),
        ("RECEBIDA", "Iniciar diagnóstico", "EM_DIAGNOSTICO", "Perfil gestor ou mecânico"),
        ("EM_DIAGNOSTICO", "Enviar orçamento", "AGUARDANDO_APROVACAO", "Orçamento atual calculado"),
        ("AGUARDANDO_APROVACAO", "Aprovar", "EM_EXECUCAO", "Token válido + estoque suficiente"),
        ("AGUARDANDO_APROVACAO", "Recusar", "ORCAMENTO_RECUSADO", "Motivo obrigatório"),
        ("EM_EXECUCAO", "Finalizar", "FINALIZADA", "Perfil gestor ou mecânico"),
        ("FINALIZADA", "Entregar", "ENTREGUE", "Perfil gestor ou atendente"),
    ], [3.7, 4.0, 4.2, 4.5])
    doc_para(doc, "5.3 Invariantes observadas no código", "Heading 2")
    doc_bullets(doc, [
        "Não duplicar o mesmo serviço ou peça na OS.", "Não calcular orçamento vazio nem enviar versão antiga.",
        "Itens só são adicionados durante diagnóstico; transições incompatíveis retornam conflito.",
        "Recusa exige motivo; aprovação repetida é idempotente e decisão oposta gera conflito.",
        "Peça inativa ou saldo insuficiente impede aprovação; transação desfaz decisão e movimentações.",
        "Token público expira, pode ser revogado, é exibido somente na emissão e armazenado como SHA-256.",
        "@Version na OS e locks pessimistas no fluxo crítico ajudam a tratar concorrência.",
    ])

    doc_para(doc, "6. APIs e segurança", "Heading 1")
    doc_para(doc, "6.1 Catálogo de endpoints", "Heading 2")
    doc_table(doc, ["Área", "Rotas", "Acesso", "Finalidade"], ENDPOINTS, [2.5, 8.4, 2.2, 4.0])
    doc_para(doc, "6.2 Autenticação e autorização", "Heading 2")
    doc_bullets(doc, [
        "Sessão stateless; CSRF desabilitado para API; filtro JWT executa antes do filtro de usuário/senha.",
        "Somente /auth/login, /publico/** e documentação OpenAPI são liberados sem JWT administrativo.",
        "Perfis GESTOR, ATENDENTE e MECANICO são aplicados com @PreAuthorize nos casos de uso.",
        "Senhas persistidas com BCrypt; segredo e expiração JWT vêm de variáveis de ambiente.",
        "Token público usa 32 bytes aleatórios, Base64 URL-safe, validade configurável e revogação explícita.",
        "Erros centralizados evitam respostas inconsistentes; DTOs validam campos e IDs positivos.",
    ])
    doc_para(doc, "6.3 Validações sensíveis", "Heading 2")
    doc_table(doc, ["Dado", "Tratamento"], [
        ("CPF/CNPJ", "Remove pontuação, rejeita sequências repetidas e valida dígitos verificadores para 11/14 dígitos."),
        ("Placa", "Converte para maiúsculas, remove espaços/hífen e aceita padrão antigo ou Mercosul."),
        ("Valores/quantidades", "BigDecimal para moeda; quantidades positivas; saldo insuficiente gera regra de negócio."),
        ("Vínculos", "Veículo ativo precisa pertencer ao cliente ativo informado na abertura da OS."),
    ], [4.0, 12.4])

    doc_para(doc, "7. Persistência e transações", "Heading 1")
    doc_para(doc, "O schema é criado por quatro migrations. Tabelas principais: usuarios, clientes, veiculos, servicos, pecas, movimentacoes_estoque, ordens_servico, itens_servico_ordem, itens_peca_ordem, orcamentos, historico_status_os e acessos_publicos_os.")
    doc_para(doc, "Decisões de consistência", "Heading 2")
    doc_bullets(doc, [
        "Foreign keys e índices sustentam integridade e consultas por cliente, veículo, status e período.",
        "Versão do orçamento é única por OS; número público da OS e documento/placa normalizados são únicos.",
        "A decisão do cliente carrega a OS para atualização; peças são bloqueadas por IDs ordenados.",
        "Aprovação, baixa, movimentações e alteração para EM_EXECUCAO participam da mesma transação.",
        "Histórico é criado pelo agregado em toda transição; relatório deriva dos instantes persistidos.",
    ])
    doc_para(doc, "Migrations", "Heading 2")
    doc_table(doc, ["Versão", "Arquivo", "Conteúdo"], [
        ("V1", "V1__baseline.sql", "Clientes, veículos, serviços, peças, movimentações e constraints iniciais."),
        ("V2", "V2__perfis_acesso.sql", "Papéis de acesso e ajustes de usuários."),
        ("V3", "V3__ordens_servico.sql", "OS, itens, orçamento e histórico."),
        ("V4", "V4__acesso_publico_e_baixa_atomica.sql", "Tokens públicos, locks/índices e rastreio da baixa automática."),
    ], [2.0, 6.2, 8.2])

    doc_para(doc, "8. Qualidade e evidências", "Heading 1")
    doc_para(doc, "8.1 Resultado reproduzido", "Heading 2")
    doc_callout(doc, "Build aprovado", "Comando executado: ./mvnw -q verify. Resultado: exit code 0; 179 testes; 0 falhas/erros; cobertura JaCoCo de linhas 717/811 = 88,41%. Integração executada contra PostgreSQL real provisionado por Testcontainers.")
    doc_para(doc, "8.2 Estratégia de testes", "Heading 2")
    doc_table(doc, ["Nível", "Exemplos", "O que protege"], [
        ("Domínio", "OrdemServicoTest, PecaTest, MovimentacaoEstoqueTest", "Transições, orçamento, recusa, saldo e movimentações."),
        ("Aplicação", "ClienteServiceTest, VeiculoServiceTest, OrdemServicoServiceTest, ServicoServiceTest", "Orquestração, validações, filtros e erros."),
        ("API", "ClienteControllerTest, ServicoControllerTest, AuthControllerTest", "Contratos HTTP, validação e segurança."),
        ("Integração", "ConsultaPublicaIntegracaoTest", "Aprovação, baixa atômica, histórico, entrega e relatório com banco real."),
        ("Gate", "JaCoCo check no phase verify", "Build falha abaixo de 80% no escopo configurado."),
    ], [3.0, 6.2, 7.2])
    doc_para(doc, "8.3 Segurança de dependências", "Heading 2")
    doc_callout(doc, "Pendência de entrega", "O enunciado exige relatório ou resultado de análise de vulnerabilidades. Não foi encontrado artefato versionado. Antes do envio, executar ferramenta SCA (por exemplo OWASP Dependency-Check), revisar falsos positivos/risco, guardar relatório sob docs/security/ e referenciá-lo no README e neste PDF.", "danger")

    doc_para(doc, "9. Execução e demonstração", "Heading 1")
    doc_para(doc, "9.1 Pré-requisitos", "Heading 2")
    doc_bullets(doc, ["Docker e Docker Compose para execução completa.", "Alternativa local: Java 21 e Docker somente para PostgreSQL.", "Arquivo .env criado a partir de .env.example, sem versionar segredos."])
    doc_para(doc, "9.2 Execução", "Heading 2")
    for line in ["cp .env.example .env", "docker compose up --build", "# API: http://localhost:8080", "# Swagger: http://localhost:8080/swagger-ui.html", "# Testes: ./mvnw verify"]:
        p = doc.add_paragraph(); p.style = doc.styles["Normal"]; p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(9); set_cell_shading if False else None
    doc_para(doc, "9.3 Credenciais de demonstração", "Heading 2")
    doc_para(doc, "O DataSeeder cria usuários de demonstração a partir de variáveis de ambiente. Não publicar senhas reais no repositório ou neste PDF; durante o vídeo, mostrar o procedimento de configuração e autenticação usando valores descartáveis.")

    doc_para(doc, "10. Links DDD, vídeo e checklist", "Heading 1")
    doc_para(doc, "10.1 Quadros colaborativos", "Heading 2")
    for label, url in MIRO.items(): doc_link_line(doc, label, url)
    doc_para(doc, "Artefatos locais", "Heading 2")
    doc_bullets(doc, [
        "docs/ddd-oficina-mecanica.excalidraw — atores, subdomínios, context map e eventos.",
        "docs/fluxo-negocio-as-is-to-be.excalidraw — processo manual e fluxo automatizado.",
        "docs/fluxo-tecnico-historias.excalidraw — épicos, histórias e dependências.",
        "docs/Fluxo-api-mapeamento.md — guia técnico; revisar trechos antigos antes da entrega.",
    ])
    doc_callout(doc, "Acesso para correção", "Abrir cada link Miro em janela anônima. Exportar os quadros para PDF/PNG e versionar em docs/ddd/, porque o enunciado exige cópia no repositório além do link externo.", "warn")
    doc_para(doc, "10.2 Vídeo", "Heading 2")
    doc_callout(doc, "PENDENTE — URL do YouTube", "Inserir link completo e público do vídeo, com no máximo 15 minutos, correspondente ao commit informado: " + COMMIT, "danger")
    doc_para(doc, "Roteiro sugerido (até 15 min)", "Heading 3")
    doc_table(doc, ["Tempo", "Conteúdo"], [
        ("0:00–1:00", "Equipe, problema, proposta e commit da entrega."),
        ("1:00–3:00", "Arquitetura monolítica modular, camadas e PostgreSQL."),
        ("3:00–5:00", "Event Storming, Bounded Contexts, agregados e linguagem ubíqua."),
        ("5:00–11:00", "Swagger: login, cliente/veículo, OS, diagnóstico, itens, orçamento, acesso público, aprovação, estoque, finalização e entrega."),
        ("11:00–12:30", "Consulta pública, histórico e relatório de tempo médio."),
        ("12:30–14:00", "Testes, 88,41% de cobertura e integração Testcontainers."),
        ("14:00–15:00", "Docker, análise de vulnerabilidades e encerramento."),
    ], [3.0, 13.4])
    doc_para(doc, "10.3 Checklist final", "Heading 2")
    doc_bullets(doc, [
        "[ ] Completar nome e RM de Kayleigh.", "[ ] Inserir vídeo público e testá-lo em janela anônima.",
        "[ ] Confirmar repositório público, branch principal protegida e commit final via PR.",
        "[ ] Exportar Miro para docs/ddd/ e testar links sem autenticação.",
        "[ ] Gerar/revisar relatório de vulnerabilidades e referenciá-lo no README.",
        "[ ] Atualizar README com arquitetura, banco, credenciais demo, DDD e vídeo.",
        "[ ] Remover documentos internos desatualizados ou corrigir afirmações incompatíveis com código.",
        "[ ] Rodar ./mvnw verify no commit final; guardar evidência da cobertura.",
        "[ ] Conferir que nenhum segredo, token ou .env real foi versionado.",
        "[ ] Exportar PDF final e abrir todos os hyperlinks antes de enviar.",
    ])

    doc_para(doc, "Referências", "Heading 1")
    for label, url in REFERENCES:
        if url.startswith("http"): doc_link_line(doc, label, url)
        else: doc_para(doc, f"{label}: {url}")
    doc_para(doc, "Fontes primárias do projeto", "Heading 2")
    doc_bullets(doc, ["Código-fonte e testes no commit informado.", "README.md, pom.xml, application.yaml, Dockerfile e docker-compose.yml.", "Migrations V1–V4 e artefatos DDD do diretório docs/.", "Relatório JaCoCo gerado por ./mvnw verify em 20/10/2026."])

    # Footer
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("OficinaTrack • Tech Challenge Fase 1 • Grupo 76  |  "); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GRAY)
        add_field(p, "PAGE")

    doc.save(OUT_DOCX)


# PDF helpers
class NumberedCanvasMixin:
    pass


def pdf_styles():
    s = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("body", parent=s["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12, textColor=colors.HexColor("#1E2933"), spaceAfter=6),
        "h1": ParagraphStyle("h1", parent=s["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#19324D"), spaceBefore=12, spaceAfter=8, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=s["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, leading=15, textColor=colors.HexColor("#2E5D7B"), spaceBefore=10, spaceAfter=5, keepWithNext=True),
        "h3": ParagraphStyle("h3", parent=s["Heading3"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#19324D"), spaceBefore=7, spaceAfter=4, keepWithNext=True),
        "cover_title": ParagraphStyle("cover_title", parent=s["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#19324D"), spaceAfter=10),
        "cover_sub": ParagraphStyle("cover_sub", parent=s["BodyText"], fontName="Helvetica", fontSize=12, leading=15, alignment=TA_CENTER, textColor=colors.HexColor("#5E6872"), spaceAfter=10),
        "small": ParagraphStyle("small", parent=s["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9.2, textColor=colors.HexColor("#34404A")),
        "table": ParagraphStyle("table", parent=s["BodyText"], fontName="Helvetica", fontSize=7.4, leading=9.2, textColor=colors.HexColor("#1E2933")),
        "table_head": ParagraphStyle("table_head", parent=s["BodyText"], fontName="Helvetica-Bold", fontSize=7.5, leading=9.3, textColor=colors.white),
        "link": ParagraphStyle("link", parent=s["BodyText"], fontName="Helvetica", fontSize=8.3, leading=10.5, textColor=colors.HexColor("#2E5D7B"), spaceAfter=5),
    }


PS = pdf_styles()


def P(text, style="body"):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), PS[style])


def link_p(label, url):
    safe = escape(url, {'"': '&quot;'})
    return Paragraph(f"<b>{escape(label)}:</b> <link href=\"{safe}\" color=\"#2E5D7B\"><u>{escape(url)}</u></link>", PS["link"])


def pdf_table(headers, rows, widths, status_col=None):
    data = [[Paragraph(escape(h), PS["table_head"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(v)).replace("\n", "<br/>"), PS["table"]) for v in row])
    t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign="CENTER")
    cmds = [
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#19324D")),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#D8DEE4")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0: cmds.append(("BACKGROUND", (0,i), (-1,i), colors.HexColor("#F7F9FA")))
        if status_col is not None:
            value = str(rows[i-1][status_col])
            fill = "#EAF5EE" if value == "ATENDIDO" else "#FBEAEA" if value == "PENDENTE" else "#FFF4CC"
            cmds += [("BACKGROUND", (status_col,i), (status_col,i), colors.HexColor(fill)), ("FONTNAME", (status_col,i), (status_col,i), "Helvetica-Bold")]
    t.setStyle(TableStyle(cmds)); return t


def callout(title, text, kind="info"):
    fill, accent = ("#FBEAEA", "#9B1C1C") if kind == "danger" else ("#FFF4CC", "#B8831B") if kind == "warn" else ("#EAF1F5", "#2E5D7B")
    content = Paragraph(f"<b><font color='#{accent[1:]}'>{escape(title)}</font></b><br/>{escape(text)}", PS["body"])
    t = Table([[content]], colWidths=[16.6*cm]); t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(fill)), ("BOX", (0,0), (-1,-1), 1.0, colors.HexColor(accent)),
        ("LEFTPADDING", (0,0), (-1,-1), 9), ("RIGHTPADDING", (0,0), (-1,-1), 9),
        ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ])); return KeepTogether([t, Spacer(1, 7)])


def bullets(items):
    return ListFlowable([ListItem(P(x), leftIndent=12) for x in items], bulletType="bullet", leftIndent=18, bulletFontName="Helvetica", bulletFontSize=6, spaceAfter=5)


def on_page(canvas, doc):
    canvas.saveState()
    if doc.page > 1:
        canvas.setStrokeColor(colors.HexColor("#D8DEE4")); canvas.setLineWidth(0.4)
        canvas.line(2.2*cm, A4[1]-1.45*cm, A4[0]-2.2*cm, A4[1]-1.45*cm)
        canvas.setFont("Helvetica", 7.5); canvas.setFillColor(colors.HexColor("#5E6872"))
        canvas.drawString(2.2*cm, A4[1]-1.25*cm, "OFICINATRACK • TECH CHALLENGE FASE 1")
        canvas.drawRightString(A4[0]-2.2*cm, 1.1*cm, f"GRUPO 76  |  {doc.page}")
    canvas.restoreState()


def build_pdf():
    frame = Frame(2.2*cm, 1.7*cm, A4[0]-4.4*cm, A4[1]-3.5*cm, id="normal")
    template = PageTemplate(id="main", frames=[frame], onPage=on_page)
    pdf = BaseDocTemplate(str(OUT_PDF), pagesize=A4, title="Tech Challenge — Fase 1 — OficinaTrack — Grupo 76", author="Grupo 76", pageTemplates=[template])
    story = []
    story += [Spacer(1, 1.2*cm), Paragraph("FIAP", ParagraphStyle("fiap", parent=PS["cover_title"], fontSize=32, textColor=colors.HexColor("#9B1C1C"))),
              P("SOAT — Software Architecture", "cover_sub"), Spacer(1, .5*cm), P("TECH CHALLENGE — FASE 1", "cover_title"),
              Paragraph("<b>OFICINATRACK</b>", ParagraphStyle("prod", parent=PS["cover_title"], fontSize=21)),
              P("Sistema Integrado de Atendimento e Execução de Serviços", "cover_sub"), P("GRUPO 76", "cover_sub"), Spacer(1, .25*cm),
              pdf_table(["Integrante", "RM", "Discord"], MEMBERS, [8.4, 3.0, 5.0]), Spacer(1, .35*cm),
              Paragraph(f"<b>Linguagem e tecnologias</b><br/>{escape(TECH)}", PS["cover_sub"]), Spacer(1, .5*cm), P("São Paulo — 2026", "cover_sub"), PageBreak()]

    def h1(x): story.append(P(x, "h1"))
    def h2(x): story.append(P(x, "h2"))
    def h3(x): story.append(P(x, "h3"))
    def para(x): story.append(P(x))
    def tbl(h, r, w, s=None): story.extend([pdf_table(h,r,w,s), Spacer(1,7)])

    h1("Controle da entrega")
    tbl(["Campo", "Valor"], [("Curso / fase / grupo", "SOAT — Software Architecture / Fase 1 / Grupo 76"), ("Repositório", REPO), ("Branch de entrega", BRANCH), ("Commit de referência", COMMIT), ("Versão do documento", "1.0 — 20/10/2026"), ("Vídeo", "PENDENTE — inserir URL pública do YouTube")], [4.2, 12.2])
    story += [link_p("Repositório completo", REPO), link_p("Versão congelada pelo commit", COMMIT_URL), callout("Bloqueios antes da entrega", "Preencher nome completo e RM de Kayleigh; inserir link público do vídeo; confirmar Miro sem login; anexar resultado da análise de vulnerabilidades; confirmar branch principal protegida e versão final via PR.", "danger")]
    h1("Sumário")
    story.append(bullets(["1. Visão executiva", "2. Matriz de aderência ao enunciado", "3. Arquitetura e decisões técnicas", "4. Domain-Driven Design", "5. Fluxos e regras de negócio", "6. APIs e segurança", "7. Persistência e transações", "8. Qualidade e evidências", "9. Execução e demonstração", "10. Links DDD, roteiro do vídeo e checklist final", "Referências"]))
    story.append(PageBreak())

    h1("1. Visão executiva")
    para("O OficinaTrack é um MVP de back-end monolítico para substituir anotações e planilhas no atendimento de uma oficina mecânica. A solução concentra cadastros, ordens de serviço, diagnóstico, orçamento, autorização do cliente, consumo de peças, histórico e indicadores em uma API REST documentada.")
    h2("Problema e resultado esperado")
    story.append(bullets(["Reduzir erros de priorização e dar rastreabilidade ao status da OS.", "Evitar perda de histórico de clientes, veículos, orçamentos e serviços.", "Controlar peças e insumos com movimentações auditáveis e baixa atômica.", "Permitir aprovação do orçamento pelo cliente sem acesso administrativo.", "Proteger operações internas com JWT e perfis de acesso."]))
    h2("Escopo entregue")
    tbl(["Dimensão", "Evidência"], [("Aplicação", "88 arquivos Java organizados em nove módulos de negócio/técnicos; 37 endpoints REST."), ("Banco", "PostgreSQL com quatro migrations Flyway e constraints de integridade."), ("Qualidade", "179 testes em 19 classes; build Maven validado; cobertura de linhas de 88,41%."), ("Operação", "Dockerfile, Docker Compose, variáveis de ambiente e Swagger UI."), ("DDD", "Diagramas de domínio, AS IS/TO BE e épicos; quatro quadros Miro; mapeamento código-domínio neste PDF.")], [3.2, 13.2])
    h1("2. Matriz de aderência ao enunciado")
    para("Status baseado no código e nos artefatos locais em 20/10/2026. Itens pendentes não são apresentados como concluídos.")
    tbl(["Requisito", "Evidência", "Status", "Ação/observação"], COMPLIANCE, [3.4, 5.7, 2.2, 5.3], 2)

    h1("3. Arquitetura e decisões técnicas")
    h2("3.1 Estilo arquitetural")
    para("Monólito modular em camadas. Cada contexto reúne API, aplicação, domínio e infraestrutura. Controllers tratam HTTP; serviços de aplicação demarcam casos de uso e transações; entidades concentram invariantes; repositórios Spring Data isolam persistência. O shared contém validações, tratamento de erros e configuração OpenAPI.")
    story.append(callout("Fluxo de uma requisição", "Cliente HTTP → Controller/DTO → Service transacional → Entidade/agregado → Repository → PostgreSQL. Respostas e erros retornam contratos próprios, sem expor entidades JPA diretamente."))
    h2("3.2 Módulos")
    tbl(["Pacote", "Responsabilidade", "Camadas principais"], [("auth", "Login, JWT, usuários e roles", "api / application / domain / infra"), ("cliente", "Cadastro e validação de documento", "api / application / domain / infra / validation"), ("veiculo", "Veículo, vínculo com cliente e placa", "api / application / domain / infra"), ("catalogoservicos", "Serviços, preço e tempo estimado", "api / application / domain / infra"), ("estoque", "Peças e movimentações", "api / application / domain / infrastructure"), ("ordemservico", "Core Domain: OS, itens, orçamento e histórico", "api / application / domain / infra"), ("publico", "Credencial por OS, consulta e decisão", "api / application / domain / infra"), ("relatorios", "Indicador de tempo médio", "api / application"), ("shared", "Erros, configuração e validações comuns", "config / exception / validation")], [3.3, 7.0, 6.1])
    h2("3.3 Decisão de banco de dados")
    para("PostgreSQL foi escolhido por oferecer transações ACID, constraints, locks de linha e consultas relacionais adequadas ao fluxo. OS, itens, versões de orçamento, histórico e movimentações possuem relações fortes; consistência é mais importante que escalabilidade independente no MVP. Flyway mantém evolução reprodutível do schema. A aprovação carrega a OS e bloqueia as peças em ordem determinística, reduzindo risco de corrida e deadlock.")
    h2("3.4 Tecnologias e justificativas")
    tbl(["Tecnologia", "Papel"], [("Java 21 / Spring Boot 4.1.1", "Base da API e injeção de dependências."), ("Spring Web MVC", "Controllers REST, validação de entrada e serialização JSON."), ("Spring Data JPA / Hibernate", "Mapeamento relacional, repositories e locking."), ("Spring Security / JJWT", "Autenticação stateless, BCrypt e JWT HS256."), ("PostgreSQL / Flyway", "Persistência transacional e migrations versionadas."), ("springdoc OpenAPI", "Swagger UI e contrato navegável."), ("Docker / Compose", "Ambiente reprodutível para API e banco."), ("JUnit / Mockito / Testcontainers / JaCoCo", "Testes isolados, integração real e cobertura mínima.")], [5.0, 11.4])

    h1("4. Domain-Driven Design")
    h2("4.1 Estratégia e Bounded Contexts")
    para("O domínio central é a evolução da Ordem de Serviço. Cadastros, estoque, identidade e relatórios dão suporte. No monólito, limites são módulos de código e linguagem, não serviços distribuídos.")
    tbl(["Contexto", "Tipo", "Raiz/centro", "Elementos", "Responsabilidade"], CONTEXTS, [2.8, 1.8, 2.8, 3.3, 5.7])
    h3("Relações do Context Map")
    story.append(bullets(["Clientes → Veículos: cliente fornece identidade usada no vínculo.", "Clientes e Veículos → OS: existência, atividade e pertencimento são validados.", "Catálogo e Estoque → OS: nomes e preços são copiados para itens históricos.", "OS → Estoque: aprovação aciona baixa síncrona e transacional.", "OS → Acesso Público: token limita consulta e decisão a uma OS.", "Identidade protege rotas administrativas; Relatórios lê histórico da OS."]))
    h2("4.2 Modelo tático e relação com o código")
    tbl(["Categoria", "Implementação", "Regra principal"], DDD_ELEMENTS, [3.0, 7.2, 6.2])
    story.append(callout("Nota de fidelidade", "Os eventos abaixo descrevem fatos de negócio usados no Event Storming. Na implementação atual, a maioria aparece como mudança de estado/histórico dentro da transação, não como classes publicadas em barramento. A baixa de estoque também é síncrona no serviço de consulta pública.", "warn"))
    h2("4.3 Linguagem Ubíqua")
    tbl(["Termo", "Definição adotada"], GLOSSARY, [5.0, 11.4])
    h2("4.4 Event Storming consolidado")
    tbl(["Fluxo", "Atores", "Comandos", "Eventos/fatos", "Políticas"], EVENT_FLOWS, [2.7, 2.2, 4.0, 3.8, 5.7])

    h1("5. Fluxos e regras de negócio")
    h2("5.1 Fluxo principal da Ordem de Serviço")
    story.append(ListFlowable([ListItem(P(s), leftIndent=14) for s in MAIN_STEPS], bulletType="1", leftIndent=20, spaceAfter=6))
    h2("5.2 Máquina de estados")
    tbl(["Origem", "Comando", "Destino", "Proteção"], [("—", "Abrir OS", "RECEBIDA", "Cliente/veículo válidos"), ("RECEBIDA", "Iniciar diagnóstico", "EM_DIAGNOSTICO", "Gestor ou mecânico"), ("EM_DIAGNOSTICO", "Enviar orçamento", "AGUARDANDO_APROVACAO", "Orçamento atual calculado"), ("AGUARDANDO_APROVACAO", "Aprovar", "EM_EXECUCAO", "Token válido + estoque"), ("AGUARDANDO_APROVACAO", "Recusar", "ORCAMENTO_RECUSADO", "Motivo obrigatório"), ("EM_EXECUCAO", "Finalizar", "FINALIZADA", "Gestor ou mecânico"), ("FINALIZADA", "Entregar", "ENTREGUE", "Gestor ou atendente")], [3.7, 4.0, 4.2, 4.5])
    h2("5.3 Invariantes observadas no código")
    story.append(bullets(["Não duplicar o mesmo serviço ou peça na OS.", "Não calcular orçamento vazio nem enviar versão antiga.", "Itens só entram durante diagnóstico; transições incompatíveis retornam conflito.", "Recusa exige motivo; aprovação repetida é idempotente e decisão oposta gera conflito.", "Peça inativa ou saldo insuficiente impede aprovação; transação desfaz decisão e movimentações.", "Token público expira, pode ser revogado, é exibido somente na emissão e armazenado como SHA-256.", "@Version na OS e locks pessimistas ajudam a tratar concorrência."]))

    h1("6. APIs e segurança")
    h2("6.1 Catálogo de endpoints")
    tbl(["Área", "Rotas", "Acesso", "Finalidade"], ENDPOINTS, [2.3, 8.1, 2.1, 3.9])
    h2("6.2 Autenticação e autorização")
    story.append(bullets(["Sessão stateless; CSRF desabilitado; filtro JWT antes do filtro de usuário/senha.", "Somente login, /publico/** e OpenAPI são liberados sem JWT administrativo.", "GESTOR, ATENDENTE e MECANICO aplicados com @PreAuthorize.", "Senhas com BCrypt; segredo e expiração JWT vêm do ambiente.", "Token público usa 32 bytes aleatórios, Base64 URL-safe, validade e revogação.", "Erros centralizados e DTOs evitam contratos inconsistentes."]))
    h2("6.3 Validações sensíveis")
    tbl(["Dado", "Tratamento"], [("CPF/CNPJ", "Remove pontuação, rejeita sequências e valida dígitos verificadores."), ("Placa", "Maiúsculas, sem espaços/hífen, padrão antigo ou Mercosul."), ("Valores/quantidades", "BigDecimal; quantidades positivas; saldo insuficiente bloqueia."), ("Vínculos", "Veículo ativo precisa pertencer ao cliente ativo da OS.")], [4.0, 12.4])

    h1("7. Persistência e transações")
    para("O schema é criado por quatro migrations. Tabelas principais: usuarios, clientes, veiculos, servicos, pecas, movimentacoes_estoque, ordens_servico, itens_servico_ordem, itens_peca_ordem, orcamentos, historico_status_os e acessos_publicos_os.")
    story.append(KeepTogether([
        P("Decisões de consistência", "h2"),
        bullets(["Foreign keys e índices sustentam integridade e consultas.", "Versão do orçamento é única por OS; número da OS e documento/placa normalizados são únicos.", "Decisão carrega OS para atualização; peças são bloqueadas por IDs ordenados.", "Aprovação, baixa, movimentações e EM_EXECUCAO compartilham transação.", "Histórico é criado em cada transição; relatório deriva dos instantes persistidos."])
    ]))
    h2("Migrations")
    tbl(["Versão", "Arquivo", "Conteúdo"], [("V1", "V1__baseline.sql", "Cadastros, estoque e constraints iniciais."), ("V2", "V2__perfis_acesso.sql", "Papéis de acesso e usuários."), ("V3", "V3__ordens_servico.sql", "OS, itens, orçamento e histórico."), ("V4", "V4__acesso_publico_e_baixa_atomica.sql", "Tokens, índices e rastreio de baixa automática.")], [2.0, 6.2, 8.2])

    h1("8. Qualidade e evidências")
    h2("8.1 Resultado reproduzido")
    story.append(callout("Build aprovado", "Comando: ./mvnw -q verify. Exit code 0; 179 testes; 0 falhas/erros; JaCoCo 717/811 linhas = 88,41%. Integração executada contra PostgreSQL via Testcontainers."))
    h2("8.2 Estratégia de testes")
    tbl(["Nível", "Exemplos", "O que protege"], [("Domínio", "OrdemServicoTest, PecaTest, MovimentacaoEstoqueTest", "Transições, orçamento, saldo e movimentações."), ("Aplicação", "ClienteServiceTest, VeiculoServiceTest, OrdemServicoServiceTest, ServicoServiceTest", "Orquestração, validações, filtros e erros."), ("API", "ClienteControllerTest, ServicoControllerTest, AuthControllerTest", "Contratos HTTP, validação e segurança."), ("Integração", "ConsultaPublicaIntegracaoTest", "Aprovação, baixa atômica, entrega e relatório."), ("Gate", "JaCoCo no phase verify", "Build falha abaixo de 80% no escopo.")], [3.0, 6.2, 7.2])
    h2("8.3 Segurança de dependências")
    story.append(callout("Pendência de entrega", "O enunciado exige relatório ou resultado de análise de vulnerabilidades. Não foi encontrado artefato versionado. Executar SCA, revisar achados, guardar relatório em docs/security/ e referenciá-lo no README e neste PDF.", "danger"))

    h1("9. Execução e demonstração")
    h2("9.1 Pré-requisitos")
    story.append(bullets(["Docker e Docker Compose para execução completa.", "Alternativa: Java 21 e Docker para PostgreSQL.", ".env criado de .env.example, sem versionar segredos."]))
    h2("9.2 Execução")
    story.append(Table([[P("cp .env.example .env\ndocker compose up --build\n# API: http://localhost:8080\n# Swagger: http://localhost:8080/swagger-ui.html\n# Testes: ./mvnw verify", "small")]], colWidths=[16.4*cm], style=[("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#F4F7F9")),("BOX",(0,0),(-1,-1),.5,colors.HexColor("#D8DEE4")),("LEFTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    h2("9.3 Credenciais de demonstração")
    para("O DataSeeder cria usuários de demonstração a partir de variáveis de ambiente. Não publicar senhas reais; no vídeo, mostrar configuração e autenticação com valores descartáveis.")

    h1("10. Links DDD, vídeo e checklist")
    h2("10.1 Quadros colaborativos")
    for label,url in MIRO.items(): story.append(link_p(label,url))
    h2("Artefatos locais")
    story.append(bullets(["docs/ddd-oficina-mecanica.excalidraw — atores, subdomínios, context map e eventos.", "docs/fluxo-negocio-as-is-to-be.excalidraw — AS IS e TO BE.", "docs/fluxo-tecnico-historias.excalidraw — épicos, histórias e dependências.", "docs/Fluxo-api-mapeamento.md — guia técnico; revisar trechos antigos antes da entrega."]))
    story.append(callout("Acesso para correção", "Abrir cada Miro em janela anônima. Exportar quadros para PDF/PNG e versionar em docs/ddd/, conforme enunciado.", "warn"))
    h2("10.2 Vídeo")
    story.append(callout("PENDENTE — URL do YouTube", "Inserir link completo e público, até 15 minutos, correspondente ao commit " + COMMIT, "danger"))
    h3("Roteiro sugerido (até 15 min)")
    tbl(["Tempo", "Conteúdo"], [("0:00–1:00", "Equipe, problema, proposta e commit."), ("1:00–3:00", "Arquitetura, camadas e PostgreSQL."), ("3:00–5:00", "DDD: Event Storming, contextos e agregados."), ("5:00–11:00", "Swagger: fluxo completo da OS e estoque."), ("11:00–12:30", "Consulta pública, histórico e relatório."), ("12:30–14:00", "Testes, cobertura e Testcontainers."), ("14:00–15:00", "Docker, vulnerabilidades e encerramento.")], [3.0, 13.4])
    h2("10.3 Checklist final")
    story.append(bullets(["[ ] Completar nome e RM de Kayleigh.", "[ ] Inserir vídeo público e testá-lo anonimamente.", "[ ] Confirmar repositório público, branch protegida e commit final via PR.", "[ ] Exportar Miro para docs/ddd/ e testar links.", "[ ] Gerar e revisar relatório de vulnerabilidades.", "[ ] Atualizar README com arquitetura, banco, DDD e vídeo.", "[ ] Corrigir documentação interna incompatível com código.", "[ ] Rodar ./mvnw verify no commit final.", "[ ] Confirmar ausência de segredos.", "[ ] Abrir hyperlinks do PDF final."]))

    h1("Referências")
    for label,url in REFERENCES:
        story.append(link_p(label,url) if url.startswith("http") else P(f"{label}: {url}"))
    h2("Fontes primárias do projeto")
    story.append(bullets(["Código-fonte e testes no commit informado.", "README.md, pom.xml, application.yaml, Dockerfile e docker-compose.yml.", "Migrations V1–V4 e artefatos DDD em docs/.", "Relatório JaCoCo gerado em 20/10/2026."]))

    pdf.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(OUT_DOCX)
    print(OUT_PDF)
